"""文档解析:支持 PDF / Word(docx) / TXT / Markdown。"""
from __future__ import annotations

import hashlib
import os
from pathlib import Path
from typing import List, Dict, Any

from .schemas import DocChunk


def _doc_id_for(path: Path) -> str:
    """基于文件内容生成稳定 doc_id(便于增量去重)。"""
    h = hashlib.md5()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(8192), b""):
            h.update(chunk)
    return h.hexdigest()[:16]


def load_pdf(path: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    parts = []
    for i, page in enumerate(reader.pages):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        if text.strip():
            parts.append(f"[第{i+1}页]\n{text}")
    return "\n\n".join(parts)


def load_docx(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    parts = []
    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            parts.append(t)
    # 表格也抓一份
    for ti, table in enumerate(doc.tables):
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)
    return "\n".join(parts)


def load_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def load_md(path: Path) -> str:
    # Markdown:直接读纯文本(分块时按段落处理即可),保留原始以便检索
    return path.read_text(encoding="utf-8", errors="ignore")


LOADERS = {
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".txt": load_txt,
    ".md": load_md,
    ".markdown": load_md,
}


def detect_and_load(path: Path) -> str:
    ext = path.suffix.lower()
    loader = LOADERS.get(ext)
    if loader is None:
        raise ValueError(
            f"暂不支持的文件类型: {ext}。支持: pdf, docx, txt, md"
        )
    return loader(path)


def parse_file(path: Path) -> Dict[str, Any]:
    """解析文件,返回 doc_id/doc_name/text。"""
    text = detect_and_load(path)
    return {
        "doc_id": _doc_id_for(path),
        "doc_name": path.name,
        "text": text,
    }